"""Findings, dashboard and exports over the collected demo estate."""
import csv
import io
import json

from openpyxl import load_workbook

from waf_inventory import report
from waf_inventory.cli import main


def test_findings_cover_the_planted_issues(session):
    got = {(f["rule"], f["subject"]) for f in report.table(session, "findings")}
    expected = {
        ("no-waf", "afd-contoso / partners-contoso-com (partners.contoso.com)"),
        ("no-waf", "dc1-bigip01.corp.contoso.com / /Common/vs_legacy_http"),
        ("detection-mode", "azure: waf-agw-shop-api"),
        ("detection-mode", "f5: /Common/hr_waf"),
        ("partial-waf", "dc1-bigip01.corp.contoso.com / /Common/vs_hr_https (hr.corp.contoso.com, selfservice.corp.contoso.com)"),
        ("l4-listener", "agw-prod-weu / lst-sftp-tcp"),
        ("old-ruleset", "azure: waf-agw-unused"),
        ("legacy-waf-config", "azure: agw-dev (legacy WAF config)"),
        ("wildcard-urls", "f5: /Common/app1_waf"),
        ("signature-staging", "f5: /Common/hr_waf"),
        ("unused-policy", "f5: /Common/old_unused_waf"),
        ("shadow-api", "GET app1.corp.contoso.com/api/v1/payroll/export"),  # not in ASM allowed URLs
        ("shadow-api", "GET shop.contoso.com/api/internal/debug/config"),  # not in the OpenAPI spec
        ("unprotected-api", "POST partners.contoso.com/api/partner/orders"),
        ("unknown-host", "old.contoso.com"),
        ("dormant-api", "GET shop.contoso.com/api/v1/wishlist"),
    }
    assert expected <= got, expected - got
    # allowed by ASM / documented in the spec -> no shadow finding
    assert ("shadow-api", "GET app1.corp.contoso.com/api/v1/employees/{employee_id}") not in got
    assert ("shadow-api", "GET shop.contoso.com/api/v1/orders/{order_id}") not in got
    assert ("waf-disabled", "f5: /Common/old_unused_waf") not in got  # unattached: reported once as unused


def test_effective_policy_on_endpoints(session):
    eps = {(e["host"], e["method"], e["template"]): e for e in report.table(session, "endpoints")}
    assert eps[("shop.contoso.com", "GET", "/api/v1/orders/{order_id}")]["waf_policy"] == "waf-agw-shop-api"
    assert eps[("shop.contoso.com", "GET", "/")]["waf_policy"] == "waf-agw-shop"
    assert eps[("app1.corp.contoso.com", "GET", "/api/v1/employees/{employee_id}")]["waf_policy"] == "/Common/app1_api_waf"
    assert eps[("hr.corp.contoso.com", "GET", "/healthz")]["mode"] == "none"  # LTM policy asm disable
    assert eps[("partners.contoso.com", "POST", "/api/partner/orders")]["mode"] == "none"


def test_dashboard_pages(client, demo_db):
    c = client
    r = c.get("/")
    assert r.status_code == 200 and "applications" in r.text and "of web listeners" in r.text and "Run all now" in r.text
    for name in report.TABLES:
        r = c.get(f"/t/{name}")
        assert r.status_code == 200, name
    r = c.get("/t/endpoints", params={"host": "api.contoso.com", "kind": "api"})
    assert "/v2/customers/{customer_id}" in r.text and "shop.contoso.com" not in r.text.split("<tbody>")[1]
    shop = c.get("/apps/shop.contoso.com").text
    assert "rr-shop/api" in shop and "waf-agw-shop-api" in shop and "detection-mode" in shop  # policy findings shown too
    assert "OpenAPI" in c.get("/apps/api.contoso.com").text
    assert c.get("/apps/10.10.20.13:80").status_code == 200  # hostname-less virtual server
    assert c.get("/apps/nope.example").status_code == 404
    pid = next(r["id"] for r in report.table(__import__("waf_inventory.models", fromlist=["x"]).connect(demo_db), "policies")
               if r["policy"] == "/Common/app1_api_waf")
    r = c.get(f"/policy/{pid}")
    assert r.status_code == 200 and "/api/v1/employees/*" in r.text
    assert c.get("/t/nope").status_code == 404 and c.get("/policy/99999").status_code == 404


def test_exports_over_http(client):
    c = client
    r = c.get("/export/findings.csv", params={"severity": "high"})
    rows = list(csv.DictReader(io.StringIO(r.text)))
    assert r.headers["content-type"].startswith("text/csv") and rows and {x["severity"] for x in rows} == {"high"}
    assert json.loads(c.get("/export/applications.json").text)[0]["application"]
    wb = load_workbook(io.BytesIO(c.get("/export/inventory.xlsx").content))
    assert wb.sheetnames == list(report.TABLES) and wb["endpoints"].max_row > 40
    assert load_workbook(io.BytesIO(c.get("/export/routes.xlsx").content)).sheetnames == ["routes"]
    html = c.get("/export/report.html").text
    assert "WAF Inventory Report" in html and "fdwafcontoso" in html and "<script" not in html
    doc = c.get("/export/openapi/api.contoso.com.json").json()
    assert "/v2/customers/{customer_id}" in doc["paths"] and "patch" in doc["paths"]["/v2/customers/{customer_id}"]
    assert c.get("/export/nope.csv").status_code == 404 and c.get("/export/routes.exe").status_code == 404


def test_exports_escape_formulas():
    rows = [{"a": "=HYPERLINK(\"x\")", "b": 1, "c": "@SUM(1)"}]
    wb = load_workbook(io.BytesIO(report.to_xlsx({"t": rows})))
    assert wb["t"]["A2"].value == "'=HYPERLINK(\"x\")" and wb["t"]["B2"].value == 1
    assert report.to_csv(rows).splitlines()[1] == "\"'=HYPERLINK(\"\"x\"\")\",1,'@SUM(1)"


def test_cli_exports(demo_db, tmp_path):
    cfg = tmp_path / "c.toml"
    cfg.write_text(f'database = "{demo_db}"\n')
    assert main(["-c", str(cfg), "export", "-f", "xlsx", "-o", str(tmp_path / "all.xlsx")]) == 0
    assert main(["-c", str(cfg), "export", "-f", "csv", "-t", "routes", "-o", str(tmp_path / "r.csv")]) == 0
    assert main(["-c", str(cfg), "export", "-f", "json", "-o", str(tmp_path / "all.json")]) == 0
    assert main(["-c", str(cfg), "export", "-f", "html", "-o", str(tmp_path / "r.html")]) == 0
    assert main(["-c", str(cfg), "export", "-f", "openapi", "-o", str(tmp_path / "oas")]) == 0
    assert set(json.loads((tmp_path / "all.json").read_text())) == set(report.TABLES)
    assert (tmp_path / "oas" / "shop.contoso.com.json").exists()
    assert "route" in (tmp_path / "r.csv").read_text().splitlines()[0]


def test_failing_source_is_recorded_and_others_continue(tmp_path, monkeypatch):
    monkeypatch.delenv("F5_X_PASSWORD", raising=False)
    cfg = tmp_path / "c.toml"
    cfg.write_text(f'''database = "{tmp_path / 'x.db'}"
[[sources]]
name = "broken"
type = "f5_rest"
host = "192.0.2.1"
username = "u"
password_env = "F5_X_PASSWORD"
[[sources]]
name = "logs"
type = "logs"
path = "{tmp_path / 'l.csv'}"
''')
    (tmp_path / "l.csv").write_text("host,method,path,status,hits\nx.com,GET,/api/a,200,9\n")
    assert main(["-c", str(cfg), "collect"]) == 1
    from waf_inventory.models import connect
    s = connect(str(tmp_path / "x.db"))
    src = {r["source"]: r for r in report.table(s, "sources")}
    assert src["broken"]["status"] == "error" and "F5_X_PASSWORD" in src["broken"]["error"]
    assert src["logs"]["status"] == "ok" and src["logs"]["endpoints"] == 1
